from pathlib import Path
from typing import List, Tuple
import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from backend.app.core.logging import get_logger
from backend.app.models.domain import DocumentType
from backend.app.models.schemas import DocumentChunk, DocumentMetadata
from backend.app.services.parser.base import (
    BaseDocumentParser,
    EmptyDocumentException,
    ParsingException,
)

logger = get_logger(__name__)


class DOCXDocumentParser(BaseDocumentParser):
    def parse(self, file_path: Path, document_type: DocumentType) -> Tuple[DocumentMetadata, List[DocumentChunk]]:
        if not file_path.exists():
            raise ParsingException(f"File not found: {file_path}")

        try:
            doc = docx.Document(str(file_path))
        except Exception as e:
            raise ParsingException(f"Corrupted or invalid DOCX file: {e}")

        chunks: List[DocumentChunk] = []
        total_chars = 0
        total_paragraphs = 0
        current_section = "Overview"
        current_page = 1
        words_on_page = 0
        chunk_counter = 1

        # Helper to check page break inside paragraph
        def has_page_break(paragraph) -> bool:
            for run in paragraph.runs:
                if "w:br" in run._r.xml and 'w:type="page"' in run._r.xml:
                    return True
            return False

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            if has_page_break(para):
                current_page += 1
                words_on_page = 0

            # Detect headings
            style_name = para.style.name.lower() if para.style else ""
            if "heading" in style_name or "title" in style_name or (len(text) < 60 and text.endswith(":")):
                current_section = text

            words = len(text.split())
            words_on_page += words
            # Realistic page pacing approximation if no explicit page breaks exist
            if words_on_page > 450:
                current_page += 1
                words_on_page = words

            total_chars += len(text)
            total_paragraphs += 1

            chunk_id = f"{document_type.value}_p{current_page}_{chunk_counter:04d}"
            chunks.append(
                DocumentChunk(
                    chunk_id=chunk_id,
                    document_type=document_type,
                    page=current_page,
                    section=current_section,
                    text=text,
                    token_count=words,
                )
            )
            chunk_counter += 1

        # Also process tables (rubrics frequently use tables!)
        for table in doc.tables:
            for row in table.rows:
                cells_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                # deduplicate adjacent merged cell text
                clean_cells = []
                for cell_t in cells_text:
                    if not clean_cells or clean_cells[-1] != cell_t:
                        clean_cells.append(cell_t)

                if clean_cells:
                    row_text = " | ".join(clean_cells)
                    words = len(row_text.split())
                    words_on_page += words
                    if words_on_page > 450:
                        current_page += 1
                        words_on_page = words

                    total_chars += len(row_text)
                    total_paragraphs += 1

                    chunk_id = f"{document_type.value}_p{current_page}_{chunk_counter:04d}"
                    chunks.append(
                        DocumentChunk(
                            chunk_id=chunk_id,
                            document_type=document_type,
                            page=current_page,
                            section=f"{current_section} (Table)",
                            text=row_text,
                            token_count=words,
                        )
                    )
                    chunk_counter += 1

        if total_chars == 0:
            raise EmptyDocumentException(f"The DOCX file '{file_path.name}' contains no readable text content.")

        metadata = DocumentMetadata(
            filename=file_path.name,
            document_type=document_type,
            page_count=current_page,
            paragraph_count=total_paragraphs,
            char_count=total_chars,
        )

        logger.info(
            f"Successfully parsed DOCX '{file_path.name}': ~{current_page} pages, "
            f"{len(chunks)} chunks, {total_chars} chars."
        )
        return metadata, chunks
