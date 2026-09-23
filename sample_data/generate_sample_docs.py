import docx
from docx.shared import Inches, Pt, RGBColor
import fitz
from pathlib import Path

sample_dir = Path(__file__).resolve().parent

SPEC_TEXT = """# ASSIGNMENT SPECIFICATION
Course: CS-504 Advanced Enterprise Cloud Systems & Governance
Assignment 2: Cloud Migration Strategy & Ethical AI Governance
Weight: 40% | Format: IEEE Style (2,250 - 2,750 words)

## 1. Executive Context & Objective
Global Retail Dynamics (GRD) operates an on-premises monolithic e-commerce and inventory platform.
Students are required to design an enterprise cloud migration architecture and ethical governance proposal for GRD's transition to modern cloud infrastructure.

## 2. Mandatory Assignment Requirements
Students must adhere strictly to the following requirements:

- Requirement R1 (Architectural Options): Formulate an enterprise cloud migration strategy evaluating at least two distinct architectural options (such as containerized Kubernetes microservices versus serverless event-driven architecture).
- Requirement R2 (Comparative Trade-off Analysis): Include a comparative trade-off matrix analyzing capital expenditure versus operational expenditure, latency bounds under peak concurrency, and disaster recovery RPO/RTO.
- Requirement R3 (Ethical AI & Data Governance): Establish a comprehensive ethical AI and data governance framework addressing customer data privacy (GDPR compliance), cross-border data sovereignty, and proactive bias auditing in automated recommendation algorithms.
- Requirement R4 (Phased Roadmap): Deliver a realistic, phased 12-month implementation timeline with staged risk mitigations and explicit rollback trigger criteria.
- Requirement R5 (Formatting & References): The report must follow IEEE citation conventions, adhere to 2,500 words (+/- 10%), and cite at least 8 peer-reviewed or industry-standard sources.
"""

RUBRIC_TEXT = """# MARKING RUBRIC
Course: CS-504 Advanced Enterprise Cloud Systems & Governance

## Criterion 1: Architectural & Technical Analysis (Weight: 30%)
- High Distinction (HD): Exemplary, rigorous comparative analysis of migration patterns with comprehensive quantitative metrics and architectural justification.
- Distinction (D): Strong technical analysis evaluating multiple options with clear rationale and minor gaps in numerical trade-off modeling.
- Credit (CR): Competent overview of cloud architectures but limited depth in comparing architectural trade-offs.
- Pass (P): Superficial description of technologies without substantive architectural justification.

## Criterion 2: Ethical AI & Data Governance (Weight: 25%)
- High Distinction (HD): Thorough integration of international ethical AI standards, proactive algorithmic bias auditing, and sovereign data lifecycle controls.
- Distinction (D): Sound ethical analysis identifying key compliance requirements with actionable mitigation strategies.
- Credit (CR): General awareness of data privacy and ethical principles without detailed implementation safeguards.
- Pass (P): Rudimentary mention of privacy regulations with negligible ethical analysis.

## Criterion 3: Strategic Implementation & Risk Mitigation (Weight: 25%)
- High Distinction (HD): Sophisticated, realistic 12-month phased roadmap with proactive rollback triggers and comprehensive risk governance.
- Distinction (D): Clear, practical implementation timeline with well-identified risk categories and mitigation protocols.
- Credit (CR): Basic project timeline with generic risks and conventional mitigations.
- Pass (P): Unrealistic or vague timeline lacking concrete milestone definitions.

## Criterion 4: Academic Rigour & Professional Communication (Weight: 20%)
- High Distinction (HD): Impeccable academic prose, coherent synthesis, and flawless citation of authoritative sources adhering to IEEE style.
- Distinction (D): Well-structured, professional report with consistent academic referencing and minor formatting flaws.
- Credit (CR): Readable report with adequate referencing but inconsistent terminology or flow.
- Pass (P): Disorganized structure, informal tone, or frequent citation inaccuracies.
"""

DRAFT_TEXT = """# Enterprise Cloud Migration Strategy and Ethical AI Governance Proposal
Author: Student 1048291
Course: CS-504 Cloud Systems & Governance

## 1. Introduction and Architectural Options
Global Retail Dynamics (GRD) currently relies on a monolithic on-premise infrastructure that struggles during seasonal sales peaks. To address scalability limitations, we evaluate two modern architectural alternatives: Option A (Containerized Microservices on Kubernetes) and Option B (Serverless Event-Driven Architecture on AWS Lambda).

Option A deploys containerized services orchestrated via Amazon Elastic Kubernetes Service (EKS). This architecture provides complete environment parity between local development and cloud production, fine-grained auto-scaling, and portable container networking across availability zones.

Option B leverages managed serverless functions through AWS Lambda, Amazon EventBridge, and DynamoDB. Serverless eliminates node provisioning overhead and scales down to zero idle costs during quiet trading windows. However, serverless architectures introduce cold-start latencies of approximately 250ms for JVM-based microservices and pose operational vendor lock-in risks.

## 2. Ethical AI Framework and Data Governance
As GRD migrates customer transaction and behavioral logs to the cloud, ethical governance is paramount. We propose an Ethical Data Governance Model structured around three core pillars:
1. Data Sovereignty & GDPR Compliance: Customer records will be partitioned into regional AWS regions (eu-central-1 for European clients and ap-southeast-2 for Australasian clients) using client-side envelope encryption with AWS KMS.
2. Algorithmic Fairness & Bias Auditing: GRD's personalized recommendation engine utilizes collaborative filtering that risks perpetuating demographic consumption disparities. We mandate monthly statistical parity and disparate impact assessments using Fairlearn to audit recommendation drift.
3. Transparent Consumer Consent: Users are provided an intuitive dashboard allowing real-time data export and immediate revocation of automated profiling permissions.

## 3. Implementation Roadmap and Risk Management
The proposed cloud migration is organized into a four-phase rollout over an 8-month transition plan:
- Phase 1 (Months 1-2): Foundation setup, VPC peering, and CI/CD security pipelines.
- Phase 2 (Months 3-4): Migration of non-critical analytics and read-only catalog replicas.
- Phase 3 (Months 5-6): Database live replication and dual-write synchronisation.
- Phase 4 (Months 7-8): Cutover of transactional checkout pipelines with canary deployments.

Key risks include network routing failures and data consistency divergence during dual-write phases. We mitigate these risks through automated health checks, blue-green deployment pipelines, and database replication heartbeats.

## 4. Conclusion and References
The containerized Kubernetes approach represents the most resilient architecture for GRD's high-throughput requirements. Combined with our strict ethical data governance controls, this migration will ensure compliance, reliability, and sustained business agility.

References:
[1] J. Dean and S. Ghemawat, "MapReduce: Simplified Data Processing on Large Clusters," Commun. ACM, vol. 51, no. 1, 2008.
[2] M. Fowler, "Microservices Guide," martinfowler.com, 2014.
[3] IEEE Global Initiative on Ethics of Autonomous and Intelligent Systems, "Ethically Aligned Design," IEEE, 2019.
[4] AWS Architecture Center, "Serverless Application Lens," Amazon Web Services, 2023.
"""

def create_pdf(text, out_path):
    doc = fitz.open()
    pages = text.split("\n\n## ")
    for i, page_content in enumerate(pages):
        page = doc.new_page()
        prefix = "## " if i > 0 else ""
        full_content = prefix + page_content
        page.insert_text((50, 60), full_content, fontsize=10.5, fontname="helv")
    doc.save(str(out_path))
    doc.close()

def create_docx(text, out_path):
    doc = docx.Document()
    for line in text.split("\n"):
        if line.startswith("# "):
            doc.add_heading(line.lstrip("# ").strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line.lstrip("# ").strip(), level=2)
        elif line.strip().startswith("- "):
            doc.add_paragraph(line.strip().lstrip("- "), style='List Bullet')
        elif line.strip():
            doc.add_paragraph(line.strip())
    doc.save(str(out_path))

# Save TXT
(sample_dir / "sample_specification.txt").write_text(SPEC_TEXT)
(sample_dir / "sample_rubric.txt").write_text(RUBRIC_TEXT)
(sample_dir / "sample_student_draft.txt").write_text(DRAFT_TEXT)

# Save PDF
create_pdf(SPEC_TEXT, sample_dir / "sample_specification.pdf")
create_pdf(RUBRIC_TEXT, sample_dir / "sample_rubric.pdf")
create_pdf(DRAFT_TEXT, sample_dir / "sample_student_draft.pdf")

# Save DOCX
create_docx(SPEC_TEXT, sample_dir / "sample_specification.docx")
create_docx(RUBRIC_TEXT, sample_dir / "sample_rubric.docx")
create_docx(DRAFT_TEXT, sample_dir / "sample_student_draft.docx")

print("Generated sample TXT, PDF, and DOCX files in sample_data/.")
